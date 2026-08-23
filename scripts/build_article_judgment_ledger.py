#!/usr/bin/env python3
"""Create a complete source-bound ledger of extracted article judgments.

This is a coverage/quarantine layer. It records reusable-looking source spans
even when they are blocked from runtime by safety, quality, case or timing
rules. It never promotes quarantined spans into core knowledge.
"""
from __future__ import annotations

import argparse
import collections
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import distill_wechat_astrology as distiller


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = Path(r"C:\Users\xiajt\Downloads\(公众号数据)表格视图.xlsx")
DEFAULT_OUTPUT = ROOT / "references/knowledge-modules/distilled"


def broad_review_spans(text: str) -> list[dict[str, Any]]:
    """Capture astrology-bearing list/label sentences missed by strict rules.

    These spans are never runtime-eligible by themselves. They exist so a
    human can decide whether an implicit judgment is reusable, historical,
    promotional residue or merely rhetoric.
    """
    spans: list[dict[str, Any]] = []
    seen: set[str] = set()
    for part in distiller.clauses(text):
        if len(part) < 12 or not distiller.ASTRO_TERMS.search(part):
            continue
        if distiller.MARKETING_TERMS.search(part):
            continue
        if re.search(r"课程|报名|咨询|扫码|欢迎加入|推荐阅读|投票|调查", part):
            continue
        normalized = re.sub(r"\s+", "", part)
        if normalized in seen:
            continue
        seen.add(normalized)
        spans.append({"source_span": part[:420], "pattern": ["astrology_bearing_clause"]})
    return spans[:40]


def merge_spans(strict: list[dict[str, Any]], broad: list[dict[str, Any]]) -> list[tuple[dict[str, Any], str]]:
    strict_keys = {re.sub(r"\s+", "", str(span["source_span"])) for span in strict}
    merged: list[tuple[dict[str, Any], str]] = [(span, "strict") for span in strict]
    for span in broad:
        key = re.sub(r"\s+", "", str(span["source_span"]))
        if key not in strict_keys:
            merged.append((span, "broad_review"))
    return merged


def empty_link_kind(url: str) -> str:
    if "/s/" in url:
        return "direct_article_link"
    if "appmsgalbum" in url:
        return "album_link"
    if url:
        return "other_link"
    return "no_link"


def build(input_path: Path, output_dir: Path) -> dict[str, Any]:
    source_hash = hashlib.sha256(input_path.read_bytes()).hexdigest()
    ws = load_workbook(input_path, read_only=True, data_only=True).active
    headers = [distiller.clean(value) for value in next(ws.iter_rows(values_only=True))]
    index = {name: position for position, name in enumerate(headers)}
    coverage_lines: list[str] = []
    judgment_lines: list[str] = []
    counts: collections.Counter[str] = collections.Counter()
    reason_counts: collections.Counter[str] = collections.Counter()
    span_counts: collections.Counter[str] = collections.Counter()
    topic_counts: collections.Counter[str] = collections.Counter()
    empty_direct_urls: set[str] = set()
    empty_album_urls: set[str] = set()
    judgment_total = 0

    def cell(values: tuple[Any, ...], name: str) -> Any:
        position = index.get(name)
        return values[position] if position is not None and position < len(values) else ""

    for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        title = distiller.clean(cell(values, "标题"))
        body = distiller.clean(cell(values, "正文"))
        cleaned_body = distiller.strip_marketing(body)
        strict_spans = distiller.structured_spans(cleaned_body)
        broad_spans = broad_review_spans(cleaned_body)
        merged_spans = merge_spans(strict_spans, broad_spans)
        reason = distiller.exclusion_reason(title, body, cleaned_body, strict_spans)
        timing = bool(distiller.TIMING_TERMS.search(title + " " + cleaned_body))
        modern = bool(distiller.MODERN_TERMS.search(title + " " + cleaned_body))
        link = distiller.clean(cell(values, "链接"))
        link_kind = empty_link_kind(link)
        if not body:
            disposition = f"empty_body_{link_kind}"
            if link_kind == "direct_article_link" and link:
                empty_direct_urls.add(link)
            if link_kind == "album_link" and link:
                empty_album_urls.add(link)
        elif reason is None and timing:
            disposition = "deferred"
        elif reason is None:
            disposition = "candidate"
        elif merged_spans:
            disposition = "quarantined"
        else:
            disposition = "excluded_no_reusable_judgment"
        counts["rows_seen"] += 1
        counts[f"rows_{disposition}"] += 1
        if body:
            counts["nonempty_body"] += 1
        if merged_spans:
            counts["articles_with_extracted_judgments"] += 1
        reason_key = reason or "candidate"
        reason_counts[reason_key] += 1
        span_counts[reason_key] += len(merged_spans)
        article_id = f"AAR-{row_no:05d}"
        judgment_ids: list[str] = []
        for position, (span, extraction_mode) in enumerate(merged_spans, start=1):
            judgment_total += 1
            judgment_id = f"AJD-{row_no:05d}-{position:02d}"
            judgment_ids.append(judgment_id)
            runtime_eligible = disposition == "candidate" and extraction_mode == "strict" and not modern
            if extraction_mode == "broad_review":
                counts["broad_review_judgments"] += 1
            else:
                counts["strict_judgments"] += 1
            judgment_lines.append(json.dumps({
                "judgment_id": judgment_id,
                "article_id": article_id,
                "source_ids": ["R60"],
                "source_locator": {
                    "workbook": input_path.name,
                    "sheet": ws.title,
                    "row": row_no,
                    "title": title,
                    "url": distiller.clean(cell(values, "链接")),
                    "published_at": distiller.serialise(cell(values, "发布日期")),
                    "source_sha256": source_hash,
                },
                "source_span": span["source_span"],
                "pattern": span["pattern"],
                "extraction_mode": extraction_mode,
                "topics": distiller.topics_for(title + " " + cleaned_body),
                "article_disposition": disposition,
                "exclusion_reason": reason,
                "timing_extension": timing,
                "modern_extension": modern,
                "runtime_eligible": runtime_eligible,
                "status": "quarantined" if disposition == "quarantined" else ("review_needed" if extraction_mode == "broad_review" else ("candidate" if disposition == "candidate" else "deferred")),
                "grade_cap": "C",
                "safety_gate": "G0_sensitive_topic_gate",
                "non_judgment_use": "source-bound research ledger only; quarantined spans cannot be used for conclusions, diagnosis, event guarantees or dates",
            }, ensure_ascii=False))
            for topic in distiller.topics_for(title + " " + cleaned_body):
                topic_counts[topic] += 1
        coverage_lines.append(json.dumps({
            "article_id": article_id,
            "source_ids": ["R60"],
            "source_locator": {
                "workbook": input_path.name,
                "sheet": ws.title,
                "row": row_no,
                "title": title,
                    "url": link,
                "published_at": distiller.serialise(cell(values, "发布日期")),
                "account": distiller.clean(cell(values, "公众号名")),
                "source_sha256": source_hash,
            },
            "body_length": len(body),
            "disposition": disposition,
            "link_kind": link_kind,
            "exclusion_reason": reason,
            "judgment_ids": judgment_ids,
            "judgment_count": len(judgment_ids),
            "strict_judgment_count": len(strict_spans),
            "broad_review_judgment_count": len(merged_spans) - len(strict_spans),
            "coverage_policy": "every workbook row is accounted for; source spans are retained as candidate, deferred, quarantined or no-reusable-judgment",
        }, ensure_ascii=False))

    output_dir.mkdir(parents=True, exist_ok=True)
    coverage_path = output_dir / "article-coverage.jsonl"
    judgments_path = output_dir / "judgment-ledger.jsonl"
    coverage_path.write_text("\n".join(coverage_lines) + "\n", encoding="utf-8")
    judgments_path.write_text("\n".join(judgment_lines) + "\n", encoding="utf-8")
    manifest = {
        "schema_version": "ARTICLE-JUDGMENT-LEDGER-0.1",
        "source_id": "R60",
        "source_file": input_path.name,
        "source_sha256": source_hash,
        "coverage_store": "distilled/article-coverage.jsonl",
        "judgment_store": "distilled/judgment-ledger.jsonl",
        "counts": {
            **dict(counts),
            "judgments": judgment_total,
            "unique_empty_body_direct_article_links": len(empty_direct_urls),
            "unique_empty_body_album_links": len(empty_album_urls),
            "articles_by_exclusion_reason": dict(reason_counts),
            "judgments_by_exclusion_reason": dict(span_counts),
        },
        "status_policy": {
            "candidate": "passes current structural and safety extraction gate; still C; modern_extension rows remain non-runtime until separately verified",
            "deferred": "timing material; record only, never activate dates without explicit timing workflow",
            "quarantined": "judgment-looking span preserved for audit but blocked by safety, sensational, case, CTA, mixed-system or quality gate",
            "excluded_no_reusable_judgment": "article accounted for, but no structured condition-mechanism judgment survived extraction",
            "empty_body_direct_article_link": "正文为空但保留了直接文章链接；抽样访问触发微信验证时停止，不绕过",
            "empty_body_album_link": "正文为空但保留了合集链接；可在公开合集页按分页继续抓取",
            "empty_body_other_link": "正文为空且链接类型未列入直接文章/合集",
            "empty_body_no_link": "正文和链接均为空",
        },
        "promotion_gate": "No promotion without source locator, independent support, counterexample, Chart Facts regression and safety gate.",
    }
    (output_dir / "judgment-ledger-index.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    retrieval_path = output_dir / "retrieval-index.json"
    if retrieval_path.exists():
        retrieval = json.loads(retrieval_path.read_text(encoding="utf-8"))
        retrieval["article_judgment_ledger"] = {
            "coverage_store": "distilled/article-coverage.jsonl",
            "judgment_store": "distilled/judgment-ledger.jsonl",
            "index": "distilled/judgment-ledger-index.json",
            "counts": {"articles": counts["rows_seen"], "nonempty_articles": counts["nonempty_body"], "judgments": judgment_total, "strict_judgments": counts["strict_judgments"], "broad_review_judgments": counts["broad_review_judgments"], "empty_body_direct_article_links": counts["rows_empty_body_direct_article_link"], "empty_body_album_links": counts["rows_empty_body_album_link"], "unique_empty_body_direct_article_links": len(empty_direct_urls), "unique_empty_body_album_links": len(empty_album_urls)},
            "by_topic": dict(topic_counts),
            "policy": "ledger is source-bound; only strict non-modern candidate judgments may be considered for bounded retrieval, and every judgment remains C until promoted",
        }
        retrieval_path.write_text(json.dumps(retrieval, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"rows": counts["rows_seen"], "nonempty": counts["nonempty_body"], "judgments": judgment_total, "candidate_rows": counts["rows_candidate"], "quarantined_rows": counts["rows_quarantined"]}, ensure_ascii=False))
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.input, args.output_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
